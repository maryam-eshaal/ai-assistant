"""AI Resume ATS Score Checker
Streamlit UI + Google Gemini Flash.

Upload a resume (PDF, DOCX or TXT), optionally paste a job description,
and get an ATS score with concrete improvements.
"""

import io
import json
import os
import re
import time

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

DEFAULT_MODEL = "gemini-3.5-flash"
# Tried in order if the chosen model is overloaded (503) or unavailable.
FALLBACK_MODELS = ["gemini-3-flash-preview", "gemini-3.5-flash-lite"]
RETRIES_PER_MODEL = 3
RETRY_BASE_DELAY = 2  # seconds; doubles after each failed attempt
MAX_RESUME_CHARS = 20000
MAX_JD_CHARS = 8000
MIN_RESUME_CHARS = 150

CATEGORIES = [
    "Formatting & Structure",
    "Keywords & Skills",
    "Work Experience & Impact",
    "Education & Certifications",
    "Clarity & Grammar",
]

SYSTEM_PROMPT = """You are an expert technical recruiter and Applicant Tracking System (ATS) analyst.
Evaluate the resume text provided by the user. Be honest, specific and strict; do not inflate scores.
If a job description is provided, judge keyword match against it. Otherwise judge general ATS readiness.

Return ONLY valid JSON with exactly this structure:
{
  "overall_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "category_scores": {
    "Formatting & Structure": <0-100>,
    "Keywords & Skills": <0-100>,
    "Work Experience & Impact": <0-100>,
    "Education & Certifications": <0-100>,
    "Clarity & Grammar": <0-100>
  },
  "strengths": ["<short point>", ...],
  "missing_keywords": ["<keyword or skill>", ...],
  "improvements": [
    {"priority": "High|Medium|Low", "issue": "<what is wrong>", "fix": "<exactly how to fix it>"}
  ],
  "rewrite_examples": [
    {"original": "<weak line from the resume>", "improved": "<stronger, quantified version>"}
  ]
}
Rules: give 3-6 strengths, 5-10 improvements ordered by priority, and 2-4 rewrite examples that
use lines actually found in the resume. Never invent facts the candidate did not state; when
suggesting numbers in rewrites, use placeholders like [X%]."""


# ----------------------------------------------------------------------------
# File reading
# ----------------------------------------------------------------------------
def extract_text(file_name: str, data: bytes) -> str:
    """Extract plain text from PDF, DOCX or TXT bytes."""
    name = file_name.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password-protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        text = "\n".join(pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        text = "\n".join(parts)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# ----------------------------------------------------------------------------
# Gemini call + response handling
# ----------------------------------------------------------------------------
def parse_json_response(raw: str) -> dict:
    """Parse model output into a dict, tolerating ```json fences and stray text."""
    raw = (raw or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.IGNORECASE).strip()
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", raw, flags=re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise ValueError("The AI returned an unreadable response. Please try again.")


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def normalize_result(data: dict) -> dict:
    """Make sure every field exists and has the right type so the UI never crashes."""
    cat = data.get("category_scores") or {}
    category_scores = {name: _clamp(cat.get(name)) for name in CATEGORIES}

    overall = data.get("overall_score")
    if overall is None and category_scores:
        overall = sum(category_scores.values()) / len(category_scores)

    def str_list(key):
        items = data.get(key) or []
        return [str(i) for i in items if str(i).strip()]

    improvements = []
    for item in data.get("improvements") or []:
        if isinstance(item, dict):
            improvements.append(
                {
                    "priority": str(item.get("priority", "Medium")).title(),
                    "issue": str(item.get("issue", "")),
                    "fix": str(item.get("fix", "")),
                }
            )
        elif str(item).strip():
            improvements.append({"priority": "Medium", "issue": str(item), "fix": ""})

    rewrites = []
    for item in data.get("rewrite_examples") or []:
        if isinstance(item, dict) and (item.get("original") or item.get("improved")):
            rewrites.append(
                {"original": str(item.get("original", "")), "improved": str(item.get("improved", ""))}
            )

    return {
        "overall_score": _clamp(overall),
        "summary": str(data.get("summary", "")),
        "category_scores": category_scores,
        "strengths": str_list("strengths"),
        "missing_keywords": str_list("missing_keywords"),
        "improvements": improvements,
        "rewrite_examples": rewrites,
    }


def _error_code(exc: Exception):
    code = getattr(exc, "code", None)
    return code if isinstance(code, int) else None


def is_retryable(exc: Exception) -> bool:
    """Temporary Google-side problems: overload, rate limit, server errors."""
    if _error_code(exc) in (429, 500, 502, 503, 504):
        return True
    text = str(exc).lower()
    return any(s in text for s in ("unavailable", "overloaded", "high demand", "resource_exhausted", "timed out"))


def is_model_missing(exc: Exception) -> bool:
    return _error_code(exc) == 404 or "not found" in str(exc).lower()


def analyze_resume(resume_text: str, job_description: str, api_key: str, model: str) -> dict:
    """Run the analysis, retrying temporary errors and falling back to other Flash models."""
    client = genai.Client(api_key=api_key)
    user_prompt = f"RESUME:\n{resume_text[:MAX_RESUME_CHARS]}\n\n"
    if job_description.strip():
        user_prompt += f"JOB DESCRIPTION:\n{job_description[:MAX_JD_CHARS]}\n"
    else:
        user_prompt += "JOB DESCRIPTION: (none provided - evaluate general ATS readiness)\n"

    config = types.GenerateContentConfig(
        system_instruction=SYSTEM_PROMPT,
        response_mime_type="application/json",
    )

    models = [model] + [m for m in FALLBACK_MODELS if m != model]
    errors = {}  # model -> last error text, shown if every model fails
    for current in models:
        for attempt in range(RETRIES_PER_MODEL):
            try:
                response = client.models.generate_content(
                    model=current, contents=user_prompt, config=config
                )
                result = normalize_result(parse_json_response(response.text))
                result["model_used"] = current
                return result
            except Exception as exc:
                errors[current] = str(exc)[:160]
                if is_retryable(exc):
                    if attempt < RETRIES_PER_MODEL - 1:
                        time.sleep(RETRY_BASE_DELAY * (2 ** attempt))
                    continue  # retry, then fall through to the next model
                if is_model_missing(exc) and current != model:
                    break  # fallback model not available on this key; try the next one
                raise  # bad API key, bad request, unreadable output, etc.

    details = "; ".join(f"{m}: {e}" for m, e in errors.items())
    raise RuntimeError(
        "No Gemini model could answer right now. If the errors below say 503/UNAVAILABLE, "
        "Google is overloaded - try again in a minute. If they say 404/NOT_FOUND, a model "
        f"name is outdated - change it in the sidebar. Details: {details}"
    )


# ----------------------------------------------------------------------------
# UI helpers
# ----------------------------------------------------------------------------
def get_api_key(sidebar_value: str) -> str:
    if sidebar_value.strip():
        return sidebar_value.strip()
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass
    return os.environ.get("GEMINI_API_KEY", "")


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 65:
        return "Good"
    if score >= 50:
        return "Needs work"
    return "Poor"


def render_results(result: dict) -> None:
    score = result["overall_score"]
    if result.get("model_used"):
        st.caption(f"Analyzed with {result['model_used']}")
    left, right = st.columns([1, 2])
    with left:
        st.metric("ATS Score", f"{score}/100", score_label(score), delta_color="off")
        st.progress(score / 100)
    with right:
        st.subheader("Summary")
        st.write(result["summary"] or "No summary returned.")

    st.divider()
    st.subheader("Score breakdown")
    for name, value in result["category_scores"].items():
        c1, c2 = st.columns([3, 1])
        c1.write(name)
        c2.write(f"**{value}/100**")
        st.progress(value / 100)

    st.divider()
    col_a, col_b = st.columns(2)
    with col_a:
        st.subheader("Strengths")
        for s in result["strengths"] or ["None identified."]:
            st.markdown(f"- {s}")
    with col_b:
        st.subheader("Missing keywords")
        if result["missing_keywords"]:
            st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))
        else:
            st.write("No major keyword gaps found.")

    st.divider()
    st.subheader("Improvements")
    order = {"High": 0, "Medium": 1, "Low": 2}
    icons = {"High": "🔴", "Medium": "🟠", "Low": "🟢"}
    for item in sorted(result["improvements"], key=lambda i: order.get(i["priority"], 1)):
        icon = icons.get(item["priority"], "🟠")
        with st.expander(f"{icon} {item['priority']}: {item['issue'][:90]}"):
            st.write(f"**Issue:** {item['issue']}")
            if item["fix"]:
                st.write(f"**How to fix:** {item['fix']}")

    if result["rewrite_examples"]:
        st.divider()
        st.subheader("Rewrite examples")
        for ex in result["rewrite_examples"]:
            st.markdown(f"**Before:** {ex['original']}")
            st.markdown(f"**After:** {ex['improved']}")
            st.write("")

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


# ----------------------------------------------------------------------------
# App
# ----------------------------------------------------------------------------
def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 AI Resume ATS Checker")
    st.caption("Upload your resume and get an ATS score with practical improvements.")

    with st.sidebar:
        st.header("Settings")
        key_input = st.text_input(
            "Gemini API key",
            type="password",
            help="Leave empty if the key is set in Streamlit secrets or GEMINI_API_KEY.",
        )
        model = st.text_input("Model", value=DEFAULT_MODEL)
        st.markdown("[Get a free API key](https://aistudio.google.com/apikey)")
        st.info("Your resume is sent to Google's Gemini API for analysis and is not stored by this app.")

    uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional, improves keyword matching)", height=150
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        api_key = get_api_key(key_input)
        if not api_key:
            st.error("Please add your Gemini API key in the sidebar or in Streamlit secrets.")
            st.stop()

        try:
            with st.spinner("Reading your resume..."):
                text = extract_text(uploaded.name, uploaded.getvalue())
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            st.stop()

        if len(text) < MIN_RESUME_CHARS:
            st.error(
                "Very little text was found. If your resume is a scanned image, "
                "ATS systems cannot read it either - export a text-based PDF or DOCX."
            )
            st.stop()

        try:
            with st.spinner("Analyzing with Gemini..."):
                st.session_state["result"] = analyze_resume(
                    text, job_description, api_key, model.strip() or DEFAULT_MODEL
                )
        except Exception as exc:
            st.error(f"Analysis failed: {exc}")
            st.stop()

    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
